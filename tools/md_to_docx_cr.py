"""
Convert docs/CHANGE_REQUEST_HTTPS_LOGIN.md → docs/CHANGE_REQUEST_HTTPS_LOGIN.docx.

Hand-written parser tuned to the markdown subset used in that document:
  - # / ## / ### / #### headings
  - GitHub-flavoured tables (| ... | ... |)
  - bullet lists (- ...) and numbered lists (1. ...) with task-checkboxes ([ ])
  - fenced code blocks (``` / ```cmd / ```js / ```sql)
  - inline **bold**, *italic*, `code`
  - blockquotes (> ...)
  - horizontal rules (---)

Run:
    python tools/md_to_docx_cr.py
"""
from __future__ import annotations
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement

REPO = Path(__file__).resolve().parent.parent

# Default targets (overridable via CLI: python md_to_docx_cr.py <src.md> [<dst.docx>])
SRC  = REPO / "docs" / "CHANGE_REQUEST_HTTPS_LOGIN.md"
DST  = REPO / "docs" / "CHANGE_REQUEST_HTTPS_LOGIN.docx"
if len(sys.argv) >= 2:
    SRC = Path(sys.argv[1]).resolve()
    DST = Path(sys.argv[2]).resolve() if len(sys.argv) >= 3 else SRC.with_suffix(".docx")

# ─── Styling constants ────────────────────────────────────────────────────────
FONT_BODY  = "Calibri"
FONT_MONO  = "Consolas"
SIZE_BODY  = Pt(10.5)
SIZE_H1    = Pt(20)
SIZE_H2    = Pt(14)
SIZE_H3    = Pt(12)
SIZE_H4    = Pt(11)
COLOR_NAVY = RGBColor(0x1E, 0x3A, 0x5F)
COLOR_TEAL = RGBColor(0x0D, 0x94, 0x88)
COLOR_GREY = RGBColor(0x55, 0x55, 0x55)
COLOR_CODE_BG = "F4F4F4"
COLOR_TABLE_HDR = "1E3A5F"
COLOR_TABLE_HDR_TXT = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_TABLE_ALT = "F5F7FA"

# ─── XML helpers ──────────────────────────────────────────────────────────────
def shade_cell(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)

def set_cell_borders(cell, color_hex="CCCCCC", size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), size)
        b.set(qn("w:color"), color_hex)
        borders.append(b)
    tc_pr.append(borders)

def add_hr(doc) -> None:
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BBBBBB")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

def shade_paragraph(p, fill_hex: str) -> None:
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    p_pr.append(shd)

# ─── Inline parser: **bold**, *italic*, `code` ────────────────────────────────
INLINE_RE = re.compile(
    r"(\*\*[^*\n]+\*\*)"          # **bold**
    r"|(\*[^*\n]+\*)"              # *italic*
    r"|(_[^_\n]+_)"                # _italic_
    r"|(`[^`\n]+`)"                # `code`
    r"|(\[[^\]]+\]\([^)]+\))"      # [text](url)
)

def add_inline_runs(p, text: str) -> None:
    """Add runs to paragraph p, decoding **bold**, *italic*, `code`, [link](url)."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            run = p.add_run(text[pos:m.start()])
            run.font.name = FONT_BODY
            run.font.size = SIZE_BODY
        token = m.group()
        if token.startswith("**") and token.endswith("**"):
            run = p.add_run(token[2:-2])
            run.bold = True
            run.font.name = FONT_BODY
            run.font.size = SIZE_BODY
        elif (token.startswith("*") and token.endswith("*")) or (
            token.startswith("_") and token.endswith("_")
        ):
            run = p.add_run(token[1:-1])
            run.italic = True
            run.font.name = FONT_BODY
            run.font.size = SIZE_BODY
        elif token.startswith("`"):
            run = p.add_run(token[1:-1])
            run.font.name = FONT_MONO
            run.font.size = Pt(9.5)
        elif token.startswith("["):
            inner = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if inner:
                run = p.add_run(inner.group(1))
                run.font.name = FONT_BODY
                run.font.size = SIZE_BODY
                run.font.color.rgb = RGBColor(0x0D, 0x6E, 0xFD)
                run.font.underline = True
        pos = m.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        run.font.name = FONT_BODY
        run.font.size = SIZE_BODY

# ─── Markdown → DOCX walker ───────────────────────────────────────────────────
def render_table(doc, rows: list[list[str]]) -> None:
    """rows[0] is the header. Build a styled table."""
    n_rows = len(rows)
    n_cols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=n_rows, cols=n_cols)
    tbl.autofit = True
    tbl.style = "Table Grid"
    for r_ix, row in enumerate(rows):
        cells = tbl.rows[r_ix].cells
        for c_ix in range(n_cols):
            text = row[c_ix] if c_ix < len(row) else ""
            cell = cells[c_ix]
            # clear default empty paragraph
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline_runs(p, text)
            set_cell_borders(cell, "CCCCCC", "4")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if r_ix == 0:
                shade_cell(cell, COLOR_TABLE_HDR)
                for run in p.runs:
                    run.font.color.rgb = COLOR_TABLE_HDR_TXT
                    run.bold = True
            elif r_ix % 2 == 0:
                shade_cell(cell, COLOR_TABLE_ALT)
    # spacing after table
    doc.add_paragraph()

def add_code_block(doc, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.2)
        shade_paragraph(p, COLOR_CODE_BG)
        run = p.add_run(line if line else " ")
        run.font.name = FONT_MONO
        run.font.size = Pt(9)
    doc.add_paragraph()  # spacer

def add_heading(doc, level: int, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
    p.paragraph_format.space_after = Pt(4)
    if level == 1:
        size, bold, color = SIZE_H1, True, COLOR_NAVY
    elif level == 2:
        size, bold, color = SIZE_H2, True, COLOR_NAVY
    elif level == 3:
        size, bold, color = SIZE_H3, True, COLOR_TEAL
    else:
        size, bold, color = SIZE_H4, True, COLOR_GREY
    run = p.add_run(text)
    run.bold = bold
    run.font.size = size
    run.font.name = FONT_BODY
    run.font.color.rgb = color
    # use heading styles for outline / TOC support
    style_name = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}[level]
    p.style = doc.styles[style_name]
    # then re-apply our run formatting because Heading style overrides font
    for r in p.runs:
        r.font.name = FONT_BODY
        r.font.size = size
        r.bold = bold
        r.font.color.rgb = color

def add_blockquote(doc, lines: list[str]) -> None:
    """Render each line as an indented + shaded paragraph."""
    for line in lines:
        text = line.lstrip("> ").rstrip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.space_after = Pt(2)
        shade_paragraph(p, "FAFAFA")
        # left border to look like a quote bar
        p_pr = p._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "6")
        left.set(qn("w:color"), "0D9488")
        p_bdr.append(left)
        p_pr.append(p_bdr)
        if text:
            add_inline_runs(p, text)
        else:
            p.add_run("")

CHECKBOX_RE = re.compile(r"^\s*-\s*\[([ xX])\]\s+(.*)$")
BULLET_RE   = re.compile(r"^(\s*)[-*]\s+(.*)$")
NUMBER_RE   = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")
HEADING_RE  = re.compile(r"^(#{1,6})\s+(.*)$")
TABLE_SEP_RE = re.compile(r"^\s*\|\s*[-:|\s]+\s*\|?\s*$")

def parse_table_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]

def add_list_item(doc, text: str, level: int, numbered: bool, checkbox: str | None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    if checkbox is not None:
        marker = "☒ " if checkbox.lower() == "x" else "☐ "
        run = p.add_run(marker)
        run.font.name = FONT_BODY
        run.font.size = SIZE_BODY
    elif numbered:
        # We don't know the ordinal here; the caller knows, but for simplicity
        # we use the heading-prefix bullets as plain bullets.
        run = p.add_run("• ")
        run.font.name = FONT_BODY
        run.font.size = SIZE_BODY
    else:
        run = p.add_run("• ")
        run.font.name = FONT_BODY
        run.font.size = SIZE_BODY
    add_inline_runs(p, text)

def add_numbered_item(doc, text: str, ordinal: int, level: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(f"{ordinal}. ")
    run.font.name = FONT_BODY
    run.font.size = SIZE_BODY
    run.bold = True
    add_inline_runs(p, text)

def parse_and_render(md_text: str, doc) -> None:
    lines = md_text.splitlines()
    i = 0
    n = len(lines)
    numbered_counters: dict[int, int] = {}  # level → next ordinal

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # blank
        if not stripped:
            numbered_counters.clear()
            i += 1
            continue

        # horizontal rule
        if stripped == "---":
            add_hr(doc)
            i += 1
            continue

        # heading
        m = HEADING_RE.match(stripped)
        if m:
            level = len(m.group(1))
            add_heading(doc, level, m.group(2).strip())
            numbered_counters.clear()
            i += 1
            continue

        # fenced code block (may be indented when nested inside a list item)
        if stripped.startswith("```"):
            i += 1
            block: list[str] = []
            while i < n and not lines[i].lstrip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1  # skip closing ```
            add_code_block(doc, block)
            continue

        # blockquote
        if stripped.startswith(">"):
            block: list[str] = []
            while i < n and (lines[i].lstrip().startswith(">") or not lines[i].strip()):
                if not lines[i].strip():
                    break
                block.append(lines[i])
                i += 1
            add_blockquote(doc, block)
            continue

        # table — at least 2 lines, line[i+1] is separator
        if stripped.startswith("|") and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1]):
            rows: list[list[str]] = [parse_table_row(lines[i])]
            i += 2  # skip header + separator
            while i < n and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            render_table(doc, rows)
            continue

        # task list / bullet / numbered
        m_chk = CHECKBOX_RE.match(line)
        if m_chk:
            checked = m_chk.group(1)
            text = m_chk.group(2)
            indent = len(line) - len(line.lstrip())
            level = indent // 2
            add_list_item(doc, text, level=level, numbered=False, checkbox=checked)
            i += 1
            continue

        m_bul = BULLET_RE.match(line)
        if m_bul:
            indent = len(m_bul.group(1))
            level = indent // 2
            add_list_item(doc, m_bul.group(2), level=level, numbered=False, checkbox=None)
            i += 1
            continue

        m_num = NUMBER_RE.match(line)
        if m_num:
            indent = len(m_num.group(1))
            level = indent // 2
            ordinal = numbered_counters.get(level, 0) + 1
            numbered_counters[level] = ordinal
            add_numbered_item(doc, m_num.group(3), ordinal, level)
            i += 1
            continue

        # plain paragraph — collect contiguous non-special lines
        para: list[str] = [line.rstrip()]
        i += 1
        while i < n:
            nxt = lines[i]
            if not nxt.strip():
                break
            if HEADING_RE.match(nxt.strip()):
                break
            if nxt.strip().startswith(("|", ">", "```", "- ", "* ")) or NUMBER_RE.match(nxt) or CHECKBOX_RE.match(nxt):
                break
            if nxt.strip() == "---":
                break
            para.append(nxt.rstrip())
            i += 1
        text = " ".join(p for p in para)
        # collapse 2+ spaces
        text = re.sub(r"\s{2,}", " ", text).strip()
        if text:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            add_inline_runs(p, text)

def main() -> int:
    if not SRC.exists():
        print(f"ERROR: source not found: {SRC}", file=sys.stderr)
        return 1

    md_text = SRC.read_text(encoding="utf-8")

    doc = Document()
    # page setup — A4 portrait, 1.6 cm margins
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width  = Cm(21.0)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Cm(1.6))

    # default font
    style = doc.styles["Normal"]
    style.font.name = FONT_BODY
    style.font.size = SIZE_BODY

    # footer with page numbers
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Endeavour Mining IT — CH-273 — page ")
    run.font.size = Pt(8)
    run.font.color.rgb = COLOR_GREY
    # PAGE field
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    fp._p.append(fld)

    parse_and_render(md_text, doc)

    DST.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DST)
    size_kb = DST.stat().st_size / 1024
    print(f"OK  wrote {DST}  ({size_kb:.1f} KB)")
    return 0

if __name__ == "__main__":
    sys.exit(main())
